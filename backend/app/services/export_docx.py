from collections import Counter
from datetime import datetime, time, timedelta
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from fastapi import HTTPException

from app.models import Thread
from app.models.base import utc_now
from app.services.catalog import get_record
from app.services.export_data import (
    METRICS,
    PLATFORMS,
    STAGES,
    TAIPEI,
    atomic_write,
    dataset,
    instant,
    permalink,
    root_path,
    segment,
)


def document(brand, title, session=None):
    from app.services.branding import DEFAULT_NAME, identity

    workspace_name = identity(session)["name"] if session is not None else DEFAULT_NAME
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft JhengHei"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.color.rgb = RGBColor.from_string("243C38")
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Title"].font.color.rgb = RGBColor.from_string("000000")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    header = section.header.paragraphs[0]
    header.text = f"{workspace_name}｜{brand}"
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 頁")
    doc.add_heading(title, 0)
    doc.add_paragraph(f"{brand}｜匯出時間 {instant(utc_now())}（台北）")

    doc.core_properties.author = workspace_name
    doc.core_properties.title = title
    return doc


def table(doc, headings, rows, widths):
    item = doc.add_table(rows=1, cols=len(headings))
    item.style = "Table Grid"
    item.autofit = False
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for key, value in {"val": "single", "sz": "4", "color": "D6DEDB"}.items():
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    item._tbl.tblPr.append(borders)
    for col, width in zip(item.columns, widths, strict=True):
        col.width = Cm(width)
    for cell, text, width in zip(item.rows[0].cells, headings, widths, strict=True):
        cell.text, cell.width = str(text), Cm(width)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    item.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, text, width in zip(item.add_row().cells, values, widths, strict=True):
            cell.text, cell.width = str(text), Cm(width)
    for index, row in enumerate(item.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement(f"w:{side}")
                edge.set(qn("w:w"), "90")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            cell._tc.get_or_add_tcPr().append(margins)
            if index == 0:
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "EAF0ED")
                cell._tc.get_or_add_tcPr().append(fill)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                if len(headings) > 2:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if len(rows) == 1 or index == 0:
                    paragraph.paragraph_format.keep_with_next = index == 0
    return item


def thread_section(doc, value):
    doc.add_heading(f"{value['kind']} #{value['id']} · {value['contact']}", 1)
    table(
        doc,
        ("項目", "內容"),
        [
            (
                "平台／帳號",
                f"{PLATFORMS.get(value['platform'], value['platform'])}／{value['account']}",
            ),
            ("回覆狀態", value["status"]),
            ("成交階段／名單", f"{value['stage']}／{value['color']}"),
            ("預約日期／數量", f"{value['booking_date']}／{value['booking_units'] or '未設定'}"),
        ],
        [4, 13],
    )
    doc.add_paragraph("人工標記已回覆只記錄處理狀態，不代表有同步到平台回覆內容。")
    if not value["messages"]:
        doc.add_paragraph("尚無訊息。")
    for message in value["messages"]:
        heading = doc.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.add_run(
            f"{message['time']}｜{message['author']}" + ("（已編輯）" if message["edited"] else "")
        ).bold = True
        lines = message["text"].splitlines()
        for index, line in enumerate(lines):
            paragraph = doc.add_paragraph(line)
            if len(message["text"]) < 800 and len(lines) <= 4:
                paragraph.paragraph_format.keep_with_next = index < len(lines) - 1
                paragraph.paragraph_format.keep_together = True
        if message["attachments"]:
            doc.add_paragraph("附件：" + "、".join(message["attachments"]) + "；請至原平台查看。")


def save(doc, root, relative):
    output = BytesIO()
    doc.save(output)
    return atomic_write(root, relative, output.getvalue())


def conversation(access, settings, thread_id):
    thread = get_record(access, Thread, thread_id)
    if thread.kind != "comment_thread":
        raise HTTPException(404, "找不到留言串")
    data = dataset(access.session, settings, thread.brand_id)
    root = root_path(access.session, settings)
    brand = data.brands[thread.brand_id]
    value = data.thread_data(data.threads[thread_id])
    access.session.commit()
    doc = document(data.clean(brand.name), "對話紀錄", access.session)
    thread_section(doc, value)
    if thread.post_id in data.posts:
        post = data.posts[thread.post_id]
        doc.add_heading("關聯貼文", 1)
        doc.add_paragraph(data.clean(post.caption))
        doc.add_paragraph(permalink(post.permalink, data.clean) or "原始連結尚未提供")
    stamp = utc_now().astimezone(TAIPEI).strftime("%Y%m%d_%H%M%S_%f")
    relative = (
        f"{segment(data.clean(brand.slug))}/users/{value['platform']}/對話_{thread_id}_{stamp}.docx"
    )
    return save(doc, root, relative)


def report(access, settings, body):
    brand = access.require_brand(body.brand_id)
    if body.end < body.start or (body.end - body.start).days > 366:
        raise HTTPException(422, "期間須依序填寫，且不得超過 367 天")
    data = dataset(access.session, settings, brand.id)
    root = root_path(access.session, settings)
    access.session.commit()
    start = datetime.combine(body.start, time.min, TAIPEI)
    end = datetime.combine(body.end + timedelta(days=1), time.min, TAIPEI)
    messages = [m for rows in data.messages.values() for m in rows if start <= m.sent_at < end]
    active = {m.thread_id for m in messages}
    stages = Counter(data.threads[t].stage for t in active)
    posts = [p for p in data.posts.values() if start <= p.published_at < end]
    doc = document(data.clean(brand.name), "期間營運報表", access.session)
    doc.add_paragraph(f"統計期間：{body.start} 至 {body.end}，包含起訖日（台北時間）。")
    doc.add_heading("訊息與成交進度", 1)
    table(
        doc,
        ("指標", "數值"),
        [
            ("新客人訊息", sum(m.direction == "in" for m in messages)),
            ("平台同步品牌回覆", sum(m.direction == "out" for m in messages)),
            ("期間內有訊息的對話", len(active)),
            ("期間發佈貼文", len(posts)),
            *((f"目前階段：{label}", stages.get(key, 0)) for key, label in STAGES.items()),
        ],
        [12, 5],
    )
    doc.add_paragraph(
        "階段為上述活躍對話的目前狀態，不是期間內的階段變更次數。"
        "人工標記 LINE 已回覆不計為平台同步回覆。"
    )
    doc.add_heading("貼文與最新數據", 1)
    doc.add_paragraph(
        "僅列期間發佈的貼文；數據為匯出當下最新快照，不代表此期間增加量。缺值標示尚無資料。"
    )
    if not posts:
        doc.add_paragraph("此期間沒有貼文。")
    for post in posts:
        account = data.accounts[post.account_id]
        metric = data.metrics.get(post.id)
        doc.add_heading(
            f"{instant(post.published_at)}｜{PLATFORMS.get(account.platform, account.platform)}", 2
        )
        caption = doc.add_paragraph(data.clean(post.caption) or "（無文字內容）")
        caption.paragraph_format.keep_with_next = True
        link = doc.add_paragraph(permalink(post.permalink, data.clean) or "原始連結尚未提供")
        link.paragraph_format.keep_with_next = True
        table(
            doc,
            list(METRICS.values()),
            [
                [
                    getattr(metric, key)
                    if metric and getattr(metric, key) is not None
                    else "尚無資料"
                    for key in METRICS
                ]
            ],
            [17 / len(METRICS)] * len(METRICS),
        )
        doc.add_paragraph("快照時間：" + (instant(metric.captured_at) if metric else "尚無資料"))
    stamp = utc_now().astimezone(TAIPEI).strftime("%Y%m%d_%H%M%S_%f")
    return save(
        doc,
        root,
        f"{segment(data.clean(brand.slug))}/reports/{body.start}_{body.end}_報表_{stamp}.docx",
    )
