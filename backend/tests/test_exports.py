import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from pydantic import SecretStr
from test_classify import classified_data as classified_data
from test_workspace import workspace as workspace

from app.config import REPO_ROOT
from app.models import Contact, Message, MetricSnapshot, Post, Staff, StaffBrand, Thread
from app.services import export_data, export_md


@pytest.fixture
def exports(client, workspace, settings):
    factory = client.app.state.session_factory
    secret = "export-must-hide-test-secret-0123456789"
    settings.gemini_api_key = SecretStr(secret)
    with factory.begin() as db:
        post = Post(
            account_id=workspace["fb_id"],
            external_id="post/unsafe:id",
            caption="測試貼文 " + secret,
            permalink="https://example.com/posts/123?access_token=hidden",
            media_type="image",
            published_at=datetime(2026, 9, 25, 16, 30, tzinfo=UTC),
        )
        db.add(post)
        db.flush()
        thread = db.get(Thread, workspace["second_id"])
        thread.post_id = post.id
        db.add(MetricSnapshot(post_id=post.id, captured_at=post.published_at, likes=0, reach=None))
        contact = db.get(Contact, db.get(Thread, workspace["thread_id"]).contact_id)
        contact.display_name = "王先生 / CON : 😃"
        message = db.get(Message, workspace["message_id"])
        message.text = "第一版內容 " + secret
        message.contact_id = contact.id
        ids = {**workspace, "post_id": post.id, "contact_id": contact.id, "secret": secret}
    return ids


def markdowns(root):
    return {str(p.relative_to(root)): p.read_text(encoding="utf-8") for p in root.rglob("*.md")}


def test_obsidian_months_incremental_metrics_atomic_and_restart(
    client, settings, exports, monkeypatch
):
    from app.services.obsidian import ObsidianWorker

    factory = client.app.state.session_factory
    with factory.begin() as db:
        thread = db.get(Thread, exports["second_id"])
        parent = Message(
            thread_id=thread.id,
            external_id="obsidian-parent",
            direction="in",
            author_type="fan",
            text="九月留言",
            sent_at=datetime(2026, 9, 20, tzinfo=UTC),
        )
        db.add(parent)
        db.flush()
        db.add(
            Message(
                thread_id=thread.id,
                external_id="obsidian-child",
                direction="in",
                author_type="fan",
                text="縮排回覆",
                sent_at=datetime(2026, 9, 21, tzinfo=UTC),
                comment_parent_known=True,
                comment_parent_id=parent.external_id,
            )
        )
        db.add(
            Message(
                thread_id=thread.id,
                external_id="obsidian-oct",
                direction="in",
                author_type="fan",
                text="十月留言",
                sent_at=datetime(2026, 10, 1, tzinfo=UTC),
            )
        )
    worker = client.app.state.obsidian_worker
    assert worker.tick(now=0)["written"] == 0
    assert worker.tick(now=19)["written"] == 0
    assert worker.tick(now=20)["ok"]
    root = settings.obsidian_root
    posts = [p for p in root.rglob("貼文.md") if p.parent.name.endswith(f"--{exports['post_id']}")]
    assert len(posts) == 1 and "Facebook" in str(posts[0])
    assert "--" in posts[0].parent.name and ":" not in posts[0].parent.name
    september = posts[0].parent / "留言 2026-09.md"
    october = posts[0].parent / "留言 2026-10.md"
    original_oct = (october.read_bytes(), october.stat().st_mtime_ns)
    assert "count:" in september.read_text(encoding="utf-8")
    assert "> **" in september.read_text(encoding="utf-8")
    assert "[[留言 2026-09]]" in posts[0].read_text(encoding="utf-8")
    with factory.begin() as db:
        db.add(
            Message(
                thread_id=thread.id,
                external_id="obsidian-new",
                direction="in",
                author_type="fan",
                text="新增九月留言",
                sent_at=datetime(2026, 9, 23, tzinfo=UTC),
            )
        )
    written = []
    original_write = export_md.atomic_write

    def write(root, path, content):
        written.append(path)
        return original_write(root, path, content)

    monkeypatch.setattr(export_md, "atomic_write", write)
    assert worker.tick(force=True)["written"] == 2
    assert not any("2026-10" in path for path in written)
    assert (october.read_bytes(), october.stat().st_mtime_ns) == original_oct
    before_month = september.read_bytes()
    with factory.begin() as db:
        db.add(
            MetricSnapshot(
                post_id=exports["post_id"], captured_at=datetime(2026, 10, 5, tzinfo=UTC), likes=42
            )
        )
    assert worker.tick(force=True)["written"] == 1
    assert september.read_bytes() == before_month
    assert "按讚：42" in posts[0].read_text(encoding="utf-8")
    assert ObsidianWorker(factory, settings).tick(force=True)["written"] == 0
    target = september.read_bytes()
    with factory.begin() as db:
        db.add(
            Message(
                thread_id=thread.id,
                external_id="obsidian-atomic",
                direction="in",
                author_type="fan",
                text="原子寫入驗收",
                sent_at=datetime(2026, 9, 24, tzinfo=UTC),
            )
        )

    def fail(*args):
        raise OSError("isolated disk failure")

    monkeypatch.setattr(export_data.os, "replace", fail)
    assert not worker.tick(force=True)["ok"]
    assert september.read_bytes() == target and not list(root.rglob("*.tmp"))


def test_obsidian_manual_sync_audits_after_concurrent_commit(client, exports, monkeypatch):
    from sqlalchemy import select

    from app.models import AuditLog, Setting

    factory = client.app.state.session_factory

    # 同步期間另一連線提交，請求原本的讀取快照已過期；稽核仍要寫入，不回 503。
    def tick(force=False):
        with factory.begin() as db:
            db.add(Setting(key="isolated.concurrent", value_json={"ok": True}))
        return {"ok": True}

    monkeypatch.setattr(client.app.state.obsidian_worker, "tick", tick)
    response = client.post("/api/exports/obsidian/sync")
    assert response.status_code == 200 and response.json()["ok"]
    with factory() as db:
        assert "obsidian.sync" in set(db.scalars(select(AuditLog.action)))


def test_obsidian_settings_permissions_conflicts_and_audit(client, settings, exports, tmp_path):
    from sqlalchemy import select

    from app.models import AuditLog

    initial = client.get("/api/settings/obsidian").json()
    value = {"version": initial["version"], "obsidian_root": str(tmp_path / "notes")}
    assert client.put("/api/settings/obsidian", json=value).status_code == 200
    assert client.put("/api/settings/obsidian", json=value).status_code == 409
    invalid = {"version": initial["version"] + 1, "obsidian_root": str(REPO_ROOT / "backend")}
    assert client.put("/api/settings/obsidian", json=invalid).status_code == 422
    assert client.post("/api/exports/obsidian/sync").json()["ok"]
    with client.app.state.session_factory() as db:
        actions = set(db.scalars(select(AuditLog.action)))
    assert {"obsidian.settings", "obsidian.sync"} <= actions
    client.headers.pop("x-csrf-token", None)
    assert client.put("/api/settings/obsidian", json=invalid).status_code == 403


def test_export_masks_platform_tokens_stored_in_database(client, settings, exports):
    from cryptography.fernet import Fernet

    from app.models import PlatformCredential
    from app.services.ai_keys import cipher

    settings.account_encryption_key = SecretStr(Fernet.generate_key().decode())
    token = "isolated-stored-page-token-0123456789"
    page_id = "9988776655"

    def sealed(value):
        return cipher(settings).encrypt(value.encode()).decode()

    factory = client.app.state.session_factory
    with factory.begin() as db:
        db.add(PlatformCredential(provider="META_PAGE_ACCESS_TOKEN", encrypted_key=sealed(token)))
        # 頁面 ID 屬 PLAIN，不是秘密；遮掉會破壞貼文網址。
        db.add(PlatformCredential(provider="META_PAGE_ID", encrypted_key=sealed(page_id)))
        db.get(Message, exports["message_id"]).text = f"觀眾貼上 {token} 與頁面 {page_id}"
    with factory() as db:
        text = export_data.ExportText(settings, db)(db.get(Message, exports["message_id"]).text)
    assert token not in text and "[憑證已遮蔽]" in text and page_id in text


def test_md_content_debounce_edits_and_restart(client, settings, exports):
    worker = client.app.state.export_worker
    first = worker.tick(now=0)
    assert first["ok"] and first["pending"] > 0 and first["written"] == 0
    assert worker.tick(now=19)["written"] == 0
    assert worker.tick(now=20)["written"] >= 3
    files = markdowns(settings.export_root)
    post_name = next(p for p, content in files.items() if "測試貼文" in content and "posts" in p)
    assert "2026-09-26_" in post_name and "unsafe_id" in post_name
    assert 'platform: "facebook"' in files[post_name]
    assert "| 按讚 | 0 |" in files[post_name] and "| 觸及 | 尚無資料 |" in files[post_name]
    assert "access_token" not in files[post_name]
    user_name = next(p for p, content in files.items() if "第一版內容" in content)
    assert "憑證已遮蔽" in files[user_name] and not any(
        exports["secret"] in v for v in files.values()
    )
    with client.app.state.session_factory.begin() as db:
        db.get(Message, exports["message_id"]).text = "第二版內容"
        db.get(Thread, exports["thread_id"]).stage = "booked"
    assert worker.tick(now=25)["written"] == 0
    assert worker.tick(now=45)["written"] == 1
    assert "第二版內容" in markdowns(settings.export_root)[user_name]
    assert "已預約" in markdowns(settings.export_root)[user_name]
    restarted = export_md.ExportWorker(client.app.state.session_factory, settings)
    assert restarted.tick(now=0)["written"] == 0
    assert restarted.status()["files"] == len(files)


def test_root_change_rebuilds_and_rename_cleans_only_owned_files(
    client, settings, exports, tmp_path
):
    worker = client.app.state.export_worker
    assert worker.tick(force=True)["ok"]
    root = settings.export_root
    keep = root / "my-notes.md"
    keep.write_text("人工筆記", encoding="utf-8")
    old = next(p for p, content in markdowns(root).items() if "第一版內容" in content)
    with client.app.state.session_factory.begin() as db:
        db.get(Contact, exports["contact_id"]).display_name = "新名稱"
    assert worker.tick(force=True)["ok"]
    assert not (root / old).exists() and keep.read_text(encoding="utf-8") == "人工筆記"
    new_root = tmp_path / "other-export"
    initial = client.get("/api/settings/export").json()
    saved = client.put(
        "/api/settings/export", json={"version": initial["version"], "export_root": str(new_root)}
    )
    assert saved.status_code == 200
    assert worker.tick()["written"] > 0
    assert len(markdowns(new_root)) == len(markdowns(root)) - 1
    assert root.exists()


def test_atomic_failure_keeps_previous_file_and_removes_temporary(tmp_path, monkeypatch):
    root = tmp_path / "output"
    target = export_data.atomic_write(root, "TEST/users/line/test.md", b"first")

    def fail(*_):
        raise OSError("private path must not appear in API error")

    monkeypatch.setattr(export_data.os, "replace", fail)
    with pytest.raises(OSError):
        export_data.atomic_write(root, "TEST/users/line/test.md", b"second")
    assert target.read_bytes() == b"first" and not list(target.parent.glob("*.tmp"))


@pytest.mark.parametrize("path", ["../escape.md", "C:/escape.md", "/escape.md"])
def test_generated_path_cannot_escape_root(tmp_path, path):
    with pytest.raises(ValueError):
        export_data.destination(tmp_path, path)


def test_settings_reject_protected_directories(client, workspace, settings):
    initial = client.get("/api/settings/export").json()
    for root in [
        REPO_ROOT,
        REPO_ROOT / ".git",
        Path("<PROTECTED_ROOT>/knowledge"),
        settings.database_path.parent,
    ]:
        result = client.put(
            "/api/settings/export", json={"version": initial["version"], "export_root": str(root)}
        )
        assert result.status_code == 422


def test_ignored_data_subfolder_is_allowed_but_not_data_itself(settings):
    assert (
        export_data.validate_root(REPO_ROOT / "data" / "export", settings)
        == (REPO_ROOT / "data" / "export").resolve()
    )
    for root in [REPO_ROOT / "data", REPO_ROOT / "backend", REPO_ROOT / "datasets" / "export"]:
        with pytest.raises(ValueError):
            export_data.validate_root(root, settings)


def test_change_during_write_is_exported_next_cycle(client, settings, exports, monkeypatch):
    worker = client.app.state.export_worker
    original = export_md.atomic_write
    edited = []

    def write(root, path, content):
        result = original(root, path, content)
        if not edited and b".md" in path.encode():
            with client.app.state.session_factory.begin() as db:
                db.get(Message, exports["message_id"]).text = "匯出同時抵達的新內容"
            edited.append(True)
        return result

    monkeypatch.setattr(export_md, "atomic_write", write)
    assert worker.tick(force=True, now=0)["ok"]
    assert worker.tick(now=1)["pending"] > 0
    assert worker.tick(now=21)["written"] == 1
    assert any("匯出同時抵達的新內容" in s for s in markdowns(settings.export_root).values())


def test_docx_conversation_report_brand_role_csrf_and_no_secrets(client, settings, exports):
    path = f"/api/exports/threads/{exports['thread_id']}/docx"
    assert client.post(path, headers={"X-CSRF-Token": "wrong"}).status_code == 403
    result = client.post(path)
    assert result.status_code == 200, result.text
    doc = Document(BytesIO(result.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "第一版內容" in text and exports["secret"] not in text
    assert "行銷管理工作檯｜" in doc.sections[0].header.paragraphs[0].text
    with ZipFile(BytesIO(result.content)) as archive:
        assert b"PAGE" in archive.read("word/footer1.xml")
        assert exports["secret"].encode() not in b"".join(
            archive.read(n) for n in archive.namelist()
        )
    assert list(settings.export_root.rglob("*.docx"))
    body = {"brand_id": exports["brand_id"], "start": "2026-09-26", "end": "2026-09-26"}
    result = client.post("/api/exports/report/docx", json=body)
    assert result.status_code == 200
    report = Document(BytesIO(result.content))
    assert any("測試貼文" in p.text for p in report.paragraphs)
    assert any("不是期間內" in p.text for p in report.paragraphs)
    assert (
        client.post("/api/exports/report/docx", json={**body, "end": "2025-01-01"}).status_code
        == 422
    )
    with client.app.state.session_factory.begin() as db:
        staff = db.get(Staff, exports["staff_id"])
        staff.role = "marketing"
        db.add(StaffBrand(staff_id=staff.id, brand_id=exports["brand_id"]))
    assert client.post(f"/api/exports/threads/{exports['other_thread']}/docx").status_code == 404
    assert (
        client.post(
            "/api/exports/report/docx", json={**body, "brand_id": exports["other_brand"]}
        ).status_code
        == 404
    )
    assert client.get("/api/exports/status").status_code == 403
    assert client.post("/api/exports/sync").status_code == 403
    with client.app.state.session_factory.begin() as db:
        db.get(Staff, exports["staff_id"]).role = "viewer"
    assert client.post(path).status_code == 403


def test_manifest_failure_status_does_not_leak_paths(client, exports, monkeypatch):
    def fail(*_):
        raise OSError(exports["secret"])

    monkeypatch.setattr(export_md, "atomic_write", fail)
    value = client.post("/api/exports/sync").json()
    assert not value["ok"] and value["error"]
    assert exports["secret"] not in json.dumps(value)
